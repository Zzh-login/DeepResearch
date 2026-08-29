import hashlib
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from docx import Document

from app.research.graph import _render_markdown
from domain.research.artifacts import GeneratedArtifact
from domain.research.report_schema import ResearchReport


def safe_resolve_artifact_path(root_path: str, storage_key: str) -> Path:
    """Resolve a database storage key without allowing path traversal."""
    root = Path(root_path).resolve()
    target = (root / storage_key).resolve()
    target.relative_to(root)
    return target


class ExportService:
    def __init__(self, settings):
        self._settings = settings
        self._root = Path(
            settings.research_artifact_storage_path
        ).resolve()

    def _safe_task_dir(self, task_id):
        path = (self._root / str(task_id)).resolve()
        path.relative_to(self._root)
        path.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _source_details(source: dict) -> list[tuple[str, str]]:
        """Return stable, human-readable provenance fields for one source."""
        source_type = str(source.get("source_type", "")).strip().lower()
        metadata = source.get("metadata") or {}
        if not isinstance(metadata, dict):
            metadata = {}

        details: list[tuple[str, str]] = []
        if source_type == "knowledge" or str(source.get("source_id", "")).startswith("K"):
            details.append(("类型", "知识库文档"))
            section = str(metadata.get("section_title") or "").strip()
            if section:
                details.append(("章节", section))
            page_number = metadata.get("page_number")
            if page_number not in (None, ""):
                details.append(("页码", str(page_number)))
            document_id = str(metadata.get("document_id") or "").strip()
            if document_id:
                details.append(("文档 ID", document_id))
            chunk_id = str(metadata.get("chunk_id") or "").strip()
            if chunk_id:
                details.append(("切片 ID", chunk_id))
        else:
            details.append(("类型", "网页来源"))
            url = str(source.get("url") or "").strip()
            if url:
                details.append(("URL", url))
        return details

    @classmethod
    def _source_markdown_block(cls, source: dict) -> str:
        source_id = str(source.get("source_id", "")).strip()
        title = str(source.get("title", "未命名来源")).strip()
        location = str(
            source.get("url") or source.get("filename") or ""
        ).strip()
        line = f"- [{source_id}] {title}"
        if location:
            line += f" - {location}"
        detail_lines = [f"  - {label}：{value}" for label, value in cls._source_details(source)]
        return "\n".join([line, *detail_lines])

    def render_markdown(
        self,
        report: ResearchReport,
        sources: list[dict] | None = None,
    ) -> str:
        markdown = _render_markdown(report)
        source_blocks = []
        for source in sources or []:
            source_id = str(source.get("source_id", "")).strip()
            if not source_id:
                continue
            source_blocks.append(
                self._source_markdown_block({
                    **source,
                    "title": str(source.get("title", "未命名来源")).replace("\n", " "),
                    "filename": str(source.get("filename") or "").replace("\n", " "),
                    "url": str(source.get("url") or "").replace("\n", " "),
                })
            )

        if not source_blocks:
            return markdown
        return markdown + "\n\n## 来源\n\n" + "\n\n".join(source_blocks)

    def render_docx(
        self,
        report: ResearchReport,
        query: str,
        output_path: Path,
        *,
        scope: str = "",
        sources: list[dict] | None = None,
    ) -> None:
        document = Document()
        document.add_heading("深度研究报告", level=1)
        document.add_paragraph(f"研究问题：{query}")
        if scope:
            document.add_paragraph(f"研究范围：{scope}")
        document.add_paragraph(
            "生成时间：" + datetime.now().astimezone().isoformat(
                timespec="seconds"
            )
        )

        section_titles = {
            "summary": "摘要",
            "findings": "主要发现",
            "uncertainties": "分歧与不确定性",
            "conclusion": "结论",
        }
        kind_titles = {
            "evidence": "证据",
            "comparison": "比较",
            "synthesis": "综合",
            "limitation": "局限",
            "recommendation": "建议",
            "table": "表格",
        }

        for block in report.blocks:
            document.add_heading(
                f"{section_titles[block.section.value]} · "
                f"{kind_titles[block.kind.value]}",
                level=2,
            )

            if block.kind.value == "table":
                table = document.add_table(
                    rows=1,
                    cols=len(block.columns),
                )
                table.style = "Table Grid"

                for cell, title in zip(
                    table.rows[0].cells,
                    block.columns,
                ):
                    cell.text = title

                for row in block.rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = value
            else:
                document.add_paragraph(block.text)

            if block.citation_ids:
                document.add_paragraph(
                    "引用：" + " ".join(
                        f"[{item}]"
                        for item in block.citation_ids
                    )
                )

        if sources:
            document.add_heading("来源", level=1)
            for source in sources:
                source_id = str(source.get("source_id", ""))
                title = str(source.get("title", "未命名来源"))
                location = (
                    source.get("url")
                    or source.get("filename")
                    or ""
                )
                line = f"[{source_id}] {title}"
                if location:
                    line += f" - {location}"
                document.add_paragraph(line, style="List Bullet")
                for label, value in self._source_details(source):
                    document.add_paragraph(
                        f"{label}：{value}",
                        style="List Bullet 2",
                    )

        document.save(output_path)

    def convert_docx_to_pdf(
        self,
        docx_path: Path,
        output_dir: Path,
    ) -> Path:
        if self._settings.research_pdf_converter.lower() != "word":
            raise RuntimeError("不支持的 PDF 转换组件")

        word_binary = Path(self._settings.microsoft_word_binary)
        if not word_binary.is_file():
            raise FileNotFoundError("Microsoft Word 未安装或路径配置错误")

        pdf_path = output_dir / f"{docx_path.stem}.pdf"
        subprocess.run(
            [
                sys.executable,
                "-m",
                "app.research.word_pdf_converter",
                str(docx_path.resolve()),
                str(pdf_path.resolve()),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=self._settings.research_export_timeout_seconds,
        )

        if not pdf_path.is_file() or pdf_path.stat().st_size <= 0:
            raise RuntimeError("PDF 导出失败")

        return pdf_path

    def generate(
        self,
        job: dict,
        *,
        scope: str = "",
        sources: list[dict] | None = None,
    ) -> GeneratedArtifact:
        """Generate one immutable artifact and atomically publish it."""
        report_payload = job["report"]
        report = ResearchReport.model_validate({
            "schema_version": report_payload.get("schema_version", 1),
            "blocks": report_payload.get("blocks", []),
        })
        scope = scope or str(report_payload.get("scope", ""))
        task_dir = self._safe_task_dir(job["task_id"])
        kind = str(job["kind"])
        extensions = {
            "markdown": "md",
            "docx": "docx",
            "pdf": "pdf",
        }
        if kind not in extensions:
            raise ValueError("不支持的导出格式")

        extension = extensions[kind]
        filename = f"research-report.{extension}"
        final_path = (task_dir / filename).resolve()
        final_path.relative_to(self._root)
        temporary = task_dir / f".{job['id']}.{extension}.tmp"
        source_docx = task_dir / f".{job['id']}.source.docx"
        converted_pdf = source_docx.with_suffix(".pdf")

        try:
            if kind == "markdown":
                temporary.write_text(
                    self.render_markdown(report, sources=sources),
                    encoding="utf-8",
                )
            elif kind == "docx":
                self.render_docx(
                    report,
                    str(job["query"]),
                    temporary,
                    scope=scope,
                    sources=sources,
                )
            else:
                self.render_docx(
                    report,
                    str(job["query"]),
                    source_docx,
                    scope=scope,
                    sources=sources,
                )
                generated_pdf = self.convert_docx_to_pdf(
                    source_docx,
                    task_dir,
                )
                os.replace(generated_pdf, temporary)

            size_bytes = temporary.stat().st_size
            max_bytes = self._settings.research_export_max_size_mb * 1024 * 1024
            if size_bytes <= 0:
                raise RuntimeError("导出文件为空")
            if size_bytes > max_bytes:
                raise RuntimeError("导出文件超过大小限制")

            digest = self._sha256(temporary)
            os.replace(temporary, final_path)
            return GeneratedArtifact(
                filename=filename,
                storage_key=final_path.relative_to(self._root).as_posix(),
                sha256=digest,
                size_bytes=size_bytes,
            )
        finally:
            temporary.unlink(missing_ok=True)
            source_docx.unlink(missing_ok=True)
            converted_pdf.unlink(missing_ok=True)
